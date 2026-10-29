from openpyxl import load_workbook


def leer_listin(archivo):
    libro = load_workbook(
        archivo,
        read_only=True,
        data_only=True
    )

    hoja = libro["LISTIN"]

    filas = []

    for fila in hoja.iter_rows(values_only=True):
        filas.append(fila)

    return filas


def generar_nombre_grupo(pasajeros):
    """
    Genera un nombre identificable para el grupo
    usando los apellidos/nombres de sus integrantes.
    """

    nombres = []

    for pasajero in pasajeros:

        nombre = pasajero["nombre"].strip()

        partes = nombre.split()

        if not partes:
            continue

        # Tomamos el primer apellido/nombre identificativo
        nombres.append(partes[0])

    if not nombres:
        return "Grupo sin nombre"

    # Evitamos nombres excesivamente largos
    nombres = nombres[:3]

    return " / ".join(nombres)


def extraer_grupos(filas):

    grupos = []
    grupo_actual = None

    for fila in filas:

        nombre = fila[2]
        cedula = fila[3]
        observacion = fila[7]

        observacion_texto = str(
            observacion or ""
        ).upper()

        nombre_texto = str(
            nombre or ""
        ).strip().upper()

        # La sección de cumpleañeros no forma parte
        # del listado de pasajeros.
        if nombre_texto == "CUMPLEAÑEROS":
            break

        if not nombre:
            continue

        nombre = str(nombre).strip()

        # Detectamos el comienzo de un nuevo grupo
        if "PUESTO" in observacion_texto:

            grupo_actual = {
                "nombre": "",
                "pasajeros": []
            }

            grupos.append(grupo_actual)

        if grupo_actual is None:
            continue

        cedula_limpia = str(
            cedula or ""
        ).strip()

        # Valores que significan que no existe cédula
        if cedula_limpia in [
            "-",
            "--",
            "---",
            "----",
            "-----",
            "------",
            "-------"
        ]:
            cedula_limpia = ""

        pasajero = {
            "nombre": nombre,
            "cedula": cedula_limpia
        }

        grupo_actual["pasajeros"].append(
            pasajero
        )

        # Actualizamos el nombre del grupo
        grupo_actual["nombre"] = generar_nombre_grupo(
            grupo_actual["pasajeros"]
        )

    # Numeramos grupos que hayan quedado sin nombre
    for indice, grupo in enumerate(grupos, start=1):

        if not grupo["nombre"]:
            grupo["nombre"] = f"Grupo {indice}"

    return grupos